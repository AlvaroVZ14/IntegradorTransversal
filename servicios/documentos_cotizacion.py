"""Word con docxtpl y PDF nativo, ambos desde el snapshot de la versión emitida."""
from io import BytesIO
from datetime import date
from html import escape
from copy import deepcopy
import re
import unicodedata

from mente_futuro.domain import code, money,quote_total
from mente_futuro.presentation import LOGO
from servicios.cotizacion_asistida import SECCIONES, EMISOR


def documento_descarga(cotizacion,config):
    """Completa documentos históricos en memoria; nunca altera la versión guardada."""
    q=deepcopy(cotizacion)
    d=q.setdefault('documento',{})
    if not d.get('emisor'):d['emisor']=deepcopy(config.get('emisor',{}))
    if 'marca' not in d:
        d['marca']={k:deepcopy(config[k]) for k in ('nombre','sitio_web','correo_contacto','logo_base64') if config.get(k)}
    if 'condiciones' not in q:q['condiciones']=config.get('condiciones','')
    return q


def nombre_archivo(cotizacion,extension='pdf'):
    cliente=str(cotizacion.get('snapshot',{}).get('cliente') or 'Cliente')
    cliente=unicodedata.normalize('NFKD',cliente).encode('ascii','ignore').decode()
    cliente=re.sub(r'[^A-Za-z0-9_-]+','_',cliente).strip('_')[:75] or 'Cliente'
    return f"Cotizacion_{code(cotizacion,'COT')}_{cliente}.{extension}"


def contenido(cotizacion):
    d=cotizacion.get('documento') or {}
    s={**cotizacion['snapshot'],**{k:d[k] for k in ('modalidad','participantes','horas') if k in d}}
    emisor=d.get('emisor',{})
    contacto=[('Cliente',s.get('cliente','')),('Contacto',s.get('contacto','')),('Correo',s.get('email','')),
        ('Teléfono',s.get('telefono','')),('Servicio',s.get('servicio','')),('Temática',s.get('tematica','')),
        ('Modalidad',s.get('modalidad','')),('Participantes',s.get('participantes','')),('Horas',f"{float(s.get('horas',0)):g}".replace('.',',')),
        ('Lugar',s.get('lugar',''))]
    bloques=[('Detalle del servicio',cotizacion['detalle'])]
    for k,label in SECCIONES.items():
        if d.get('mostrar',{}).get(k,True) and d.get('textos',{}).get(k):bloques.append((label,d['textos'][k]))
    descuento=f"{float(cotizacion['descuento_pct']):g}".replace('.',',')+'%'
    motivo=cotizacion.get('descuento_motivo') or d.get('descuento_motivo')
    if motivo:descuento+=' · '+motivo
    base=int(cotizacion['base_clp']); porcentaje=float(cotizacion['descuento_pct'])
    neto=quote_total(base,cotizacion['descuento_pct'],0)
    importes=[('Valor base',money(base)),('Descuento',descuento),('Neto después de descuento',money(neto)),
        ('Impuesto','Exento de IVA' if cotizacion['iva_pct']==0 else 'IVA 19%'),('Total CLP',money(cotizacion['total_clp']))]
    return dict(codigo=code(cotizacion,'COT'),version=str(cotizacion.get('version_cotizacion',1)),
        emision=date.fromisoformat(cotizacion['fecha_emision']).strftime('%d/%m/%Y'),
        vigencia=date.fromisoformat(cotizacion['valida_hasta']).strftime('%d/%m/%Y'),
        emisor=[(label,str(emisor.get(k,''))) for k,label in EMISOR.items() if emisor.get(k)],
        contacto=[(k,str(v or 'Por definir')) for k,v in contacto],bloques=bloques,importes=importes,
        condiciones=cotizacion.get('condiciones',''),marca=d.get('marca') or {},
        servicio=s.get('servicio') or 'Servicio',participantes=str(s.get('participantes') or 'Por definir'),
        horas=f"{float(s.get('horas',0)):g}".replace('.',','),modalidad=s.get('modalidad') or 'Por definir',
        detalle=cotizacion['detalle'],base=money(cotizacion['base_clp']),total=money(cotizacion['total_clp']))


def generar_word(cotizacion):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docxtpl import DocxTemplate
    datos=contenido(cotizacion)
    doc=Document()
    sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.65)
    sec.left_margin=sec.right_margin=Inches(.8)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(7)
    for n in (1,2):
        estilo=doc.styles['Heading '+str(n)];estilo.font.name='Arial'
        estilo.font.color.rgb=RGBColor.from_string('3D6FE0');estilo.font.size=Pt(14 if n==1 else 12)
        estilo.paragraph_format.keep_with_next=True
    doc.styles['Title'].font.color.rgb=RGBColor(0,0,0);doc.styles['Title'].font.size=Pt(24)
    # La plantilla base de python-docx trae una regla bajo Title; se elimina.
    for borde in doc.styles['Title'].element.xpath('.//w:pBdr'):borde.getparent().remove(borde)
    doc.add_picture(str(LOGO),width=Inches(1.65))
    doc.add_paragraph('Cotización de servicios',style='Title')
    doc.add_paragraph('{{ codigo }} · Versión {{ version }}')
    doc.add_paragraph('Emisión {{ emision }} · Vigente hasta {{ vigencia }}')
    if datos['emisor']:
        doc.add_heading('Emisor',level=1)
        doc.add_paragraph('{{ emisor_texto }}')
    doc.add_heading('Datos de la propuesta',level=1)
    tabla=doc.add_table(rows=0,cols=2);tabla.style='Light Shading Accent 1'
    contexto={k:datos[k] for k in ('codigo','version','emision','vigencia')}
    contexto['emisor_texto']='\n'.join(f'{k}: {v}' for k,v in datos['emisor'])
    for i,(label,valor) in enumerate(datos['contacto']):
        celdas=tabla.add_row().cells;celdas[0].text=label;celdas[1].text='{{ contacto_'+str(i)+' }}';contexto['contacto_'+str(i)]=valor
    for i,(label,valor) in enumerate(datos['bloques']):
        doc.add_heading(label,level=1);doc.add_paragraph('{{ bloque_'+str(i)+' }}');contexto['bloque_'+str(i)]=valor
    doc.add_heading('Inversión',level=1)
    tabla=doc.add_table(rows=0,cols=2);tabla.style='Light Shading Accent 1'
    for i,(label,valor) in enumerate(datos['importes']):
        celdas=tabla.add_row().cells;celdas[0].text=label;celdas[1].text='{{ importe_'+str(i)+' }}';contexto['importe_'+str(i)]=valor
    doc.add_heading('Condiciones comerciales',level=1)
    doc.add_paragraph('{{ condiciones }}');contexto['condiciones']=datos['condiciones']
    doc.core_properties.title='Cotización de servicios '+datos['codigo'];doc.core_properties.author='Mente y Futuro'
    plantilla=BytesIO();doc.save(plantilla);plantilla.seek(0)
    tpl=DocxTemplate(plantilla);tpl.render(contexto,autoescape=True)
    salida=BytesIO();tpl.save(salida)
    return salida.getvalue()


def generar_pdf(cotizacion):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,HRFlowable
    datos=contenido(cotizacion);salida=BytesIO()
    ancho,alto=A4;margen=40;util=ancho-2*margen
    oscuro=colors.HexColor('#0B2545');amarillo=colors.HexColor('#FFCC00')
    gris=colors.HexColor('#F6F8FC');borde=colors.HexColor('#E4E9F2')
    estilos=getSampleStyleSheet()
    estilos.add(ParagraphStyle(name='Mfy',fontName='Helvetica',fontSize=11,leading=15,spaceAfter=5,textColor=oscuro,splitLongWords=True))
    estilos.add(ParagraphStyle(name='MfyHeading',parent=estilos['Mfy'],fontName='Helvetica-Bold',fontSize=11,spaceBefore=13,spaceAfter=5,keepWithNext=True))
    estilos.add(ParagraphStyle(name='MfySmall',parent=estilos['Mfy'],fontSize=9,leading=12))
    estilos.add(ParagraphStyle(name='MfyWhite',parent=estilos['Mfy'],fontName='Helvetica-Bold',textColor=colors.white))
    estilos.add(ParagraphStyle(name='MfyAmount',parent=estilos['Mfy'],alignment=2))
    def p(texto,estilo='Mfy'):
        return Paragraph(escape(str(texto)).replace('\n','<br/>'),estilos[estilo])
    def parrafos(texto):
        # Cada párrafo puede continuar en la página siguiente sin arrastrar
        # una sección entera y dejar huecos grandes en la página anterior.
        return [p(linea) if linea.strip() else Spacer(1,5) for linea in str(texto).splitlines()]
    def seccion(titulo):
        linea=HRFlowable(width='100%',thickness=.7,color=amarillo,spaceAfter=7)
        linea.keepWithNext=True
        return [p(titulo.upper(),'MfyHeading'),linea]
    def tabla(filas,widths,encabezado=False):
        t=Table(filas,colWidths=widths,hAlign='LEFT',repeatRows=1 if encabezado else 0,splitByRow=1,splitInRow=1)
        reglas=[('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.4,borde),
            ('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),
            ('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),5)]
        if encabezado:reglas.extend([('BACKGROUND',(0,0),(-1,0),oscuro),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,gris])])
        else:reglas.append(('ROWBACKGROUNDS',(0,0),(-1,-1),[gris,colors.white]))
        t.setStyle(TableStyle(reglas))
        return t
    marca=datos['marca'];emisor=dict(datos['emisor'])
    nombre=marca.get('nombre') or emisor.get('Razón social','')
    # El logo de marca instalado es el respaldo para versiones anteriores.
    logo=str(LOGO)
    if marca.get('logo_base64'):
        import base64
        from PIL import Image as PILImage
        encoded=marca['logo_base64'].split(',',1)[-1]
        if len(encoded)>7000000:raise ValueError('El logo configurado supera el tamaño permitido.')
        raw=base64.b64decode(encoded,validate=True)
        with PILImage.open(BytesIO(raw)) as im:
            if im.format not in ('PNG','JPEG'):raise ValueError('Configura un logo PNG o JPEG.')
            im.verify()
        logo=BytesIO(raw)
    imagen=Image(logo,width=120,height=66.5,kind='proportional',hAlign='LEFT')
    contacto_empresa='\n'.join(str(v) for v in (nombre,marca.get('correo_contacto') or emisor.get('Correo del emisor'),
                                               emisor.get('Teléfono del emisor'),marca.get('sitio_web')) if v)
    cabecera=Table([[imagen,p(contacto_empresa,'MfySmall')]],colWidths=[145,util-145])
    cabecera.setStyle(TableStyle([('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),0)]))
    titulo=Table([[p('COTIZACIÓN DE SERVICIOS','MfyWhite')]],colWidths=[util])
    titulo.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),oscuro),('TOPPADDING',(0,0),(-1,-1),10),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
    historia=[cabecera,Spacer(1,12),titulo,Spacer(1,10),
        tabla([[p('N° COTIZACIÓN','MfySmall'),p('FECHA DE EMISIÓN','MfySmall'),p('VÁLIDA HASTA','MfySmall')],
               [p(datos['codigo']),p(datos['emision']),p(datos['vigencia'])]],[util/3]*3),
        p('Versión '+datos['version'],'MfySmall')]
    historia+=seccion('Dirigido a')
    # Solo datos reales existentes en el caso; no se inventan cantidades ni precios.
    contacto=[(k,v) for k,v in datos['contacto'] if k in ('Cliente','Contacto','Correo','Teléfono','Lugar')]
    historia.append(tabla([[p(k,'MfySmall'),p(v)] for k,v in contacto],[110,util-110]))
    historia+=seccion('Detalle del servicio')
    historia.append(tabla([[p(k,'MfySmall'),p(v)] for k,v in [
        ('Servicio',datos['servicio']),('Modalidad',datos['modalidad']),('Duración (horas)',datos['horas']),
        ('Participantes',datos['participantes'])]],[110,util-110]))
    historia.append(Spacer(1,9));historia.extend(parrafos(datos['detalle']))
    for label,valor in datos['bloques']:
        if label=='Detalle del servicio':continue
        historia+=seccion(label);historia.extend(parrafos(valor))
    historia+=seccion('Inversión')
    historia.append(tabla([[p('CONCEPTO','MfyWhite'),p('VALOR (CLP)','MfyWhite')],
        [p(datos['servicio']),p(datos['base'],'MfyAmount')]],[util-125,125],True))
    totales=tabla([[p(k),p(v,'MfyAmount')] for k,v in datos['importes'][1:]],[util-180,180])
    totales.setStyle(TableStyle([('BACKGROUND',(0,-1),(-1,-1),colors.HexColor('#CCCCFF')),
                                ('LINEABOVE',(0,-1),(-1,-1),1,oscuro)]))
    historia.append(totales)
    historia+=seccion('Condiciones comerciales');historia.extend(parrafos(datos['condiciones']))
    legales=[(k,v) for k,v in datos['emisor'] if k not in ('Correo del emisor','Teléfono del emisor')]
    if legales:
        historia+=seccion('Datos del emisor')
        historia.append(p('\n'.join(f'{k}: {v}' for k,v in legales),'MfySmall'))
    def pie(canvas,doc):
        canvas.setStrokeColor(borde);canvas.line(margen,34,ancho-margen,34)
        canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#5B6B82'))
        canvas.drawString(margen,22,datos['codigo']+' · Versión '+datos['version'])
        canvas.drawRightString(ancho-margen,22,'Página '+str(doc.page))
    SimpleDocTemplate(salida,pagesize=A4,leftMargin=margen,rightMargin=margen,topMargin=35,bottomMargin=48,
        title='Cotización '+datos['codigo'],author=str(nombre)).build(historia,onFirstPage=pie,onLaterPages=pie)
    return salida.getvalue()
